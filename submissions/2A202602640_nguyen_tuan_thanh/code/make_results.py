"""make_results.py - Bước 5: gom log thật thành results.xlsx, biểu đồ tổng hợp và report.md.

Mọi con số đều đọc từ logs/ (summary.json, history.csv của từng exp_id; step*_*.json) và từ eval_out/ (kết quả của
eval.py gốc trên predictions/). Không có số nào gõ tay; số trích dẫn từ nguồn khác (bài báo, bảng timm) được ghi rõ
là TRÍCH DẪN.

    build_xlsx(lab)     -> <sub>/results.xlsx  (7 sheet bắt buộc của GUIDE 6.1 + Bonus + Setup)
    write_report(lab)   -> <sub>/report.md     (bản nháp đầy đủ 9 mục, câu chữ sinh theo số liệu; các nhận xét
                           cần mắt người - nhìn ảnh bị đoán sai - được đánh dấu để rà lại)
    check_consistency(lab) -> kiểm tra exp_id nào cũng có ảnh curves/ và số trong xlsx khớp eval.py
"""
from __future__ import annotations

import json
import math
import sys
from pathlib import Path

import numpy as np
import pandas as pd

CODE_DIR = Path(__file__).resolve().parent
if str(CODE_DIR) not in sys.path:
    sys.path.insert(0, str(CODE_DIR))

from dataset import CLASS_NAMES  # noqa: E402
from experiments import AXIS_NAMES, LN9, PAPER_TABLE1  # noqa: E402

# TRÍCH DẪN (không phải kết quả của bài): top-1 ImageNet-1k của đúng các tag trọng số, theo bảng kết quả của timm
# (https://github.com/huggingface/pytorch-image-models/blob/main/results/results-imagenet.csv).
IMAGENET_TOP1_CITED = {
    "resnet50.tv_in1k": 76.13, "resnext50_32x4d.tv_in1k": 77.62, "convnext_tiny.fb_in1k": 82.07,
    "deit_small_patch16_224.fb_in1k": 79.86, "swin_tiny_patch4_window7_224.ms_in1k": 81.38,
    "efficientnet_b0.ra_in1k": 77.69, "mobilenetv3_large_100.ra_in1k": 75.77,
}
PAPER = {"resnet50_acc": 95.7, "inception_acc": 95.1, "chinee_apple": 88.5, "snake_weed": 88.8}  # TRÍCH DẪN bài báo gốc


# --------------------------------------------------------------------------- #
# Đọc log
# --------------------------------------------------------------------------- #
def all_summaries(lab) -> list[dict]:
    out = []
    for p in sorted(lab.paths.logs.glob("*/seed*/summary.json")):
        out.append(json.loads(p.read_text(encoding="utf-8")))
    return out


def history(lab, exp_id: str, seed: int) -> pd.DataFrame:
    return pd.read_csv(lab.paths.logs / exp_id / f"seed{seed}" / "history.csv")


def eval_summary(lab, tag: str) -> dict | None:
    p = lab.paths.eval_out / f"{tag}_summary.json"
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else None


def pm(mean, std, digits: int = 4) -> str:
    if std is None or (isinstance(std, float) and math.isnan(std)):
        return f"{mean:.{digits}f} (1 seed)"
    return f"{mean:.{digits}f} ± {std:.{digits}f}"


def curve_notes(h: pd.DataFrame) -> dict:
    """Mô tả đường cong của một lần chạy từ history.csv (dùng cho nhận xét hội tụ / quá khớp)."""
    f1 = h["val_macro_f1"].to_numpy()
    vl, tl = h["val_loss"].to_numpy(), h["train_loss"].to_numpy()
    best = int(f1.argmax())
    reach = int(np.argmax(f1 >= 0.99 * f1.max())) + 1
    i_min = int(vl.argmin())
    overfit = bool(i_min < len(vl) - 1 and vl[-1] > vl[i_min] * 1.10 and tl[-1] < tl[i_min])
    return {"best_epoch": best + 1, "epoch_reach_99pct": reach, "val_loss_min_epoch": i_min + 1,
            "val_loss_min": float(vl[i_min]), "val_loss_last": float(vl[-1]), "train_loss_last": float(tl[-1]),
            "overfit_flag": overfit}


# --------------------------------------------------------------------------- #
# Các bảng
# --------------------------------------------------------------------------- #
def tables(lab) -> dict[str, pd.DataFrame]:
    s1, s2, s3, s4 = (lab.load(k) for k in ("step1_backbones", "step2_training", "step3_inference", "step4_final"))
    T: dict[str, pd.DataFrame] = {}

    if s1:
        ch = s1["choice"]
        T["Backbones"] = pd.DataFrame([{
            "exp_id": r["exp_id"], "backbone": r["backbone"].split(".")[0], "tag trọng số (timm)": r["weights_tag"],
            "#tham số (M)": r["params_m"], "GMAC": r["gmacs"], "độ phân giải": r["img_size"], "epoch": r["epochs"],
            "seed": r["seed"], "macro-F1 val": r["val_macro_f1"], "top-1 val": r["val_top1"],
            "thời gian train/epoch (s)": r["train_time_per_epoch_s"], "độ trễ batch-1 p50 (ms)": r["latency_b1_p50_ms"],
            "độ trễ batch-1 p95 (ms)": r["latency_b1_p95_ms"], "epoch tốt nhất": r["best_epoch"],
            "ghi chú": "; ".join(x for x in (
                r["group"], "đóng băng backbone, chỉ train head (linear probe)" if r["init"] == "frozen" else "tinh chỉnh toàn bộ, công thức nền T00",
                "ĐƯỢC CHỌN đi tiếp" if r["exp_id"] == ch["chosen_exp_id"] else "",
                "macro-F1 val cao nhất" if r["exp_id"] == ch["best_exp_id"] else "", "1 seed") if x)}
            for r in s1["rows"]])

    if s2:
        s = s2["choice"]["noise_std_t00"]

        def verdict(r):
            if r["axis"] == "-":
                return r["note"]
            d = r["delta_vs_t00"]
            if not math.isfinite(s):
                v = "chưa có std"
            elif abs(d) <= s:
                v = f"|Δ| <= s={s:.4f}: không phân biệt được với nhiễu"
            else:
                v = f"Δ {'>' if d > 0 else '<'} {'+' if d > 0 else '-'}s={s:.4f}: {'có lợi' if d > 0 else 'có hại'} (1 seed)"
            extra = [r["note"]] if r["note"] else []
            if r["exp_id"] == s2["choice"]["final_recipe_exp_id"]:
                extra.append("CÔNG THỨC CHUNG KẾT (-> F01)")
            return "; ".join([v, *extra])

        T["Training"] = pd.DataFrame([{
            "exp_id": r["exp_id"], "backbone": r["backbone"].split(".")[0], "trục thay đổi": r["axis_name"],
            "khác T00 ở điểm nào": r["diff"], "seed": r["seed"], "macro-F1 val": r["val_macro_f1"],
            "top-1 val": r["val_top1"], "Δ macro-F1 so với T00 (seed 0)": r["delta_vs_t00"],
            "F1 val Chinee Apple": r["f1_chinee_apple"], "F1 val Snake Weed": r["f1_snake_weed"],
            "F1 val lớp cỏ thấp nhất": r["f1_min_weed"], "F1 val Negatives": r["f1_negatives"],
            "epoch tốt nhất": r["best_epoch"], "ghi chú": verdict(r)} for r in s2["rows"]])

    if s3:
        T["Inference"] = pd.DataFrame([{
            "exp_id": r["exp_id"], "phương pháp": r["method"], "mô hình/checkpoint": r.get("model"),
            "K (số view hoặc số mô hình)": r.get("k"), "macro-F1 val": r.get("val_macro_f1"),
            "top-1 val": r.get("val_top1"), "ECE val": r.get("val_ece"),
            "độ trễ p50 batch-1 (ms)": r.get("lat_p50_ms"), "độ trễ p95 batch-1 (ms)": r.get("lat_p95_ms"),
            "độ trễ p99 batch-1 (ms)": r.get("lat_p99_ms"), "thông lượng (ảnh/s; batch xem ghi chú/Latency)": r.get("throughput_img_s"),
            "chi phí tương đối so với I00": r.get("rel_cost"), "Δ macro-F1 so với I00": r.get("delta_f1_vs_i00"),
            "ghi chú": "; ".join(x for x in (r.get("note", ""),
                                             "ĐƯỢC CHỌN cho chung kết (ngoại tuyến)" if r["exp_id"] == s3["choice"]["offline_method"] else "") if x)}
            for r in s3["rows"]])
        lat = [{"cấu hình": r["config"], "GPU": r["gpu"], "dtype": r["dtype"], "batch": r["batch"],
                "kích thước ảnh": r["img_size"], "gộp BN": "có" if r["fused_bn"] else "không",
                "p50 (ms)": r["p50"], "p95 (ms)": r["p95"], "p99 (ms)": r["p99"], "ảnh/s": r["images_per_s"],
                "số lần đo": r["n"], "warmup (lần)": r["warmup"], "torch": r["torch"],
                "tính tiền xử lý": "không"} for r in s3["latency"]]
        if s1:
            lat += [{"cấu hình": f"{r['exp_id']} {r['backbone']} (1 view)", "GPU": r["latency"]["gpu"], "dtype": "fp32",
                     "batch": 1, "kích thước ảnh": r["img_size"], "gộp BN": "không", "p50 (ms)": r["latency"]["p50"],
                     "p95 (ms)": r["latency"]["p95"], "p99 (ms)": r["latency"]["p99"],
                     "ảnh/s": r["latency"]["images_per_s"], "số lần đo": r["latency"]["n"],
                     "warmup (lần)": r["latency"]["warmup"], "torch": r["latency"]["torch"], "tính tiền xử lý": "không"}
                    for r in s1["rows"]]
        T["Latency"] = pd.DataFrame(lat)

    if s4:
        ch = s4["inference"]
        bb = lab.final_cfg(lab.plan.seeds[0]).backbone
        recipe = s2["choice"]["final_recipe_exp_id"] if s2 else "?"
        desc = {
            "T00": f"MỐC: {bb} + công thức nền T00 + I00 (1 view, không scaling)",
            "F01": f"CHUNG KẾT: {bb} + công thức {recipe} + {ch['offline_method']} ({ch['offline_desc']}) + temperature scaling",
            "F01_uncal": "F01 trước temperature scaling (cùng mô hình, cùng view)",
            "F02": f"THỜI GIAN THỰC: {bb} + công thức {recipe} + I00 (1 view) + temperature scaling",
        }
        rows = []
        for tag in ("T00", "F01", "F01_uncal", "F02"):
            ps = pd.read_csv(lab.paths.eval_out / f"{tag}_per_seed.csv")
            es = eval_summary(lab, tag)
            key = "F01" if tag == "F01_uncal" else tag
            for _, r in ps.iterrows():
                seed_info = next(x for x in s4["per_seed"] if x["seed"] == int(r["seed"]))
                rows.append({"exp_id": tag, "cấu hình (backbone + công thức + suy luận)": desc[tag], "seed": int(r["seed"]),
                             "macro-F1 val": seed_info[key]["val_macro_f1"], "macro-F1 test": r["macro_f1"],
                             "top-1 test": r["top1"], "balanced acc test": r["balanced_acc"], "ECE test": r["ece"],
                             "T (khớp trên val)": {"F01": seed_info["T_offline"], "F02": seed_info["T_realtime"]}.get(tag),
                             "mean ± std qua seed": ""})
            vals = [x[key]["val_macro_f1"] for x in s4["per_seed"]]
            rows.append({"exp_id": tag, "cấu hình (backbone + công thức + suy luận)": f"TỔNG HỢP {len(ps)} seed (std mẫu, ddof=1)",
                         "seed": "mean", "macro-F1 val": float(np.mean(vals)), "macro-F1 test": es["macro_f1"]["mean"],
                         "top-1 test": es["top1"]["mean"], "balanced acc test": es["balanced_acc"]["mean"],
                         "ECE test": es["ece"]["mean"], "T (khớp trên val)": None,
                         "mean ± std qua seed": f"macro-F1 test {pm(es['macro_f1']['mean'], es['macro_f1']['std'])}; "
                                                f"top-1 test {pm(es['top1']['mean'], es['top1']['std'])}; "
                                                f"ECE test {pm(es['ece']['mean'], es['ece']['std'])}; "
                                                f"macro-F1 val {pm(float(np.mean(vals)), float(np.std(vals, ddof=1)) if len(vals) > 1 else float('nan'))}"})
            rows.append({"exp_id": tag, "cấu hình (backbone + công thức + suy luận)": "", "seed": "std",
                         "macro-F1 val": float(np.std(vals, ddof=1)) if len(vals) > 1 else None,
                         "macro-F1 test": es["macro_f1"]["std"], "top-1 test": es["top1"]["std"],
                         "balanced acc test": es["balanced_acc"]["std"], "ECE test": es["ece"]["std"],
                         "T (khớp trên val)": None, "mean ± std qua seed": ""})
        T["Final"] = pd.DataFrame(rows)

        pc_rows = []
        for tag, label in (("T00", "mốc T00 + I00"), ("F01", "chung kết F01 (tốt nhất)"), ("F02", "thời gian thực F02")):
            pc = pd.read_csv(lab.paths.eval_out / f"{tag}_per_class.csv")
            for _, r in pc.iterrows():
                pc_rows.append({"cấu hình": label, "exp_id": tag, "lớp": r["class"], "số ảnh test": int(r["support"]),
                                "precision (mean)": r["precision_mean"], "precision (std)": r["precision_std"],
                                "recall (mean)": r["recall_mean"], "recall (std)": r["recall_std"],
                                "F1 (mean)": r["f1_mean"], "F1 (std)": r["f1_std"]})
        T["PerClass"] = pd.DataFrame(pc_rows)

    # Summary: top 10 cấu hình theo macro-F1 val, kèm chi phí/độ trễ
    cand = []
    lat_of = {r["backbone"]: r["latency_b1_p50_ms"] for r in (s1["rows"] if s1 else [])}
    for s in all_summaries(lab):
        kind = {"B": "backbone", "T": "công thức huấn luyện", "F": "chung kết (1 view)", "X": "bonus"}.get(s["exp_id"][0], "?")
        cand.append({"exp_id": s["exp_id"] + ("" if s["seed"] == 0 else f" (seed {s['seed']})"), "loại": kind,
                     "mô tả": f"{s['backbone'].split('.')[0]}; {s['diff']}", "macro-F1 val": s["val_macro_f1"],
                     "top-1 val": s["val_top1"], "K": 1, "độ trễ p50 batch-1 (ms)": lat_of.get(s["backbone"]),
                     "chi phí tương đối (so với 1 view cùng backbone)": 1.0})
    if s3:
        for r in s3["rows"]:
            if r.get("val_macro_f1") is None or r["exp_id"] == "I00":
                continue
            cand.append({"exp_id": r["exp_id"], "loại": "suy luận", "mô tả": f"{r['method']} [{r['model']}]",
                         "macro-F1 val": r["val_macro_f1"], "top-1 val": r["val_top1"], "K": r["k"],
                         "độ trễ p50 batch-1 (ms)": r["lat_p50_ms"],
                         "chi phí tương đối (so với 1 view cùng backbone)": r["rel_cost"]})
    if cand:
        top = pd.DataFrame(cand).sort_values("macro-F1 val", ascending=False).head(10).reset_index(drop=True)
        top.insert(0, "hạng", range(1, len(top) + 1))
        top["ghi chú"] = "chọn trên val; " + top["loại"].map(
            lambda k: "1 seed" if k in ("backbone", "công thức huấn luyện", "suy luận", "bonus") else "mỗi dòng 1 seed")
        if s4:
            ef, et = eval_summary(lab, "F01"), eval_summary(lab, "T00")
            extra = pd.DataFrame([
                {"hạng": "TEST", "exp_id": "F01", "loại": "chung kết", "mô tả": "macro-F1 TEST " + pm(ef["macro_f1"]["mean"], ef["macro_f1"]["std"])
                 + "; top-1 TEST " + pm(ef["top1"]["mean"], ef["top1"]["std"]), "macro-F1 val": None, "ghi chú": f"{len(ef['seeds'])} seed, test 1 lần/seed"},
                {"hạng": "TEST", "exp_id": "T00", "loại": "mốc", "mô tả": "macro-F1 TEST " + pm(et["macro_f1"]["mean"], et["macro_f1"]["std"])
                 + "; top-1 TEST " + pm(et["top1"]["mean"], et["top1"]["std"]), "macro-F1 val": None, "ghi chú": f"{len(et['seeds'])} seed, test 1 lần/seed"},
            ])
            top = pd.concat([top, extra], ignore_index=True)
        T["Summary"] = top

    # Bonus + Setup
    bonus = []
    rb = lab.load("bonus_robustness")
    if rb:
        bonus += [{"phần": "Lệch phân phối (val)", "mô hình/runtime": r["model"], "điều kiện": r["corruption"],
                   "macro-F1 val": r["val_macro_f1"], "top-1 val": r["val_top1"], "ECE trước TS": r["ece_before_ts"],
                   "ECE sau TS (T từ val sạch)": r["ece_after_ts"], "recall Negatives": r["recall_negatives"],
                   "recall lớp cỏ thấp nhất": r["min_weed_recall"]} for r in rb["rows"]]
    ox = lab.load("bonus_onnx")
    if ox:
        bonus += [{"phần": "ONNX", "mô hình/runtime": r["runtime"], "điều kiện": r.get("error", "batch 1, 224"),
                   "p50 (ms)": r.get("p50"), "p95 (ms)": r.get("p95"), "p99 (ms)": r.get("p99"),
                   "lệch logit tối đa so với PyTorch": r.get("max_abs_diff_vs_pytorch")} for r in ox.get("rows", [])]
    if bonus:
        T["Bonus"] = pd.DataFrame(bonus)
    sums = all_summaries(lab)
    if sums:
        s0 = sums[0]
        setup = [("GPU/thiết bị huấn luyện", s0["device"]), *[(f"phiên bản {k}", v) for k, v in s0["versions"].items()],
                 ("fold", 0), ("seed quét sàng (B, T)", lab.plan.seeds[0]), ("seed chung kết và mốc", str(list(lab.plan.seeds))),
                 ("epoch", lab.plan.epochs), ("batch size", lab.plan.batch_size), ("công cụ đếm GMAC", s0["gmac_tool"]),
                 ("số lần huấn luyện", len(sums)),
                 ("tổng thời gian huấn luyện (giờ)", round(sum(s["total_time_s"] for s in sums) / 3600, 2))]
        T["Setup"] = pd.DataFrame(setup, columns=["mục", "giá trị"])
    return T


# --------------------------------------------------------------------------- #
# results.xlsx
# --------------------------------------------------------------------------- #
def build_xlsx(lab) -> Path:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    T = tables(lab)
    order = ["Summary", "Backbones", "Training", "Inference", "Final", "PerClass", "Latency", "Bonus", "Setup"]
    best_key = {"Backbones": "macro-F1 val", "Training": "macro-F1 val", "Inference": "macro-F1 val",
                "Latency": None, "Summary": None, "PerClass": None, "Final": None}
    path = lab.paths.sub / "results.xlsx"
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name in order:
            if name not in T:
                continue
            df = T[name]
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.freeze_panes = "A2"
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="2F4B6E")
                c.alignment = Alignment(wrap_text=True, vertical="center")
            ws.row_dimensions[1].height = 34
            for j, col in enumerate(df.columns, start=1):
                width = max([len(str(col)) * 0.6] + [len(f"{v:.4f}" if isinstance(v, float) else str(v)) for v in df[col].head(60)])
                ws.column_dimensions[get_column_letter(j)].width = float(min(70, max(9, width + 2)))
                for row in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                    v = row[0].value
                    if isinstance(v, float):
                        low = str(col).lower()
                        if any(k in low for k in ("(ms)", "ảnh/s", "(s)", "thời gian")):
                            row[0].number_format = "0.00"
                        elif any(k in low for k in ("tham số", "gmac", "chi phí")):
                            row[0].number_format = "0.00"
                        else:
                            row[0].number_format = "0.0000"
            hl = PatternFill("solid", fgColor="FFF2B3")
            rows_hl = []
            key = best_key.get(name)
            if key and key in df and df[key].notna().any():
                rows_hl.append(int(df[key].astype(float).idxmax()))
            if name == "Final":
                rows_hl += [i for i, v in enumerate(df["seed"]) if v == "mean" and df["exp_id"][i] == "F01"]
            if name == "Summary":
                rows_hl += [0] + [i for i, v in enumerate(df["hạng"]) if v == "TEST" and df["exp_id"][i] == "F01"]
            if name == "PerClass":
                rows_hl += [i for i in range(len(df)) if df["exp_id"][i] == "F01" and df["lớp"][i] in ("Chinee apple", "Snake weed")]
            for i in rows_hl:
                for c in ws[i + 2]:
                    c.fill = hl
    return path


# --------------------------------------------------------------------------- #
# Biểu đồ tổng hợp
# --------------------------------------------------------------------------- #
def backbone_figure(lab) -> Path | None:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    s1 = lab.load("step1_backbones")
    if not s1:
        return None
    fig, ax = plt.subplots(1, 2, figsize=(14, 5))
    for r in s1["rows"]:
        color = "#d9662b" if r["init"] == "frozen" else "#2a78b5"
        for a, xkey in zip(ax, ("latency_b1_p50_ms", "gmacs")):
            a.scatter(r[xkey], r["val_macro_f1"], s=30 + 9 * r["params_m"], color=color, alpha=0.75, zorder=3)
            a.annotate(f"{r['exp_id']} {r['backbone'].split('.')[0]}", (r[xkey], r["val_macro_f1"]),
                       textcoords="offset points", xytext=(6, 5), fontsize=8)
    ax[0].set(xlabel="độ trễ p50 batch 1, FP32 (ms)", ylabel="macro-F1 val", title="Macro-F1 val theo độ trễ")
    ax[1].set(xlabel="GMAC (224x224)", ylabel="macro-F1 val", title="Macro-F1 val theo GMAC")
    for a in ax:
        a.grid(alpha=0.3)
    fig.suptitle("So sánh backbone, công thức nền T00, 1 seed (cỡ chấm ~ số tham số; cam = đóng băng backbone)", fontsize=10)
    fig.tight_layout()
    path = lab.paths.figures / "backbones_tradeoff.png"
    fig.savefig(path, dpi=130)
    plt.close(fig)
    return path


def ablation_figure(lab) -> Path | None:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    s2 = lab.load("step2_training")
    if not s2:
        return None
    rows = [r for r in s2["rows"] if r["axis"] != "-"]
    s = s2["choice"]["noise_std_t00"]
    fig, ax = plt.subplots(figsize=(10, 0.42 * len(rows) + 1.8))
    ypos = np.arange(len(rows))[::-1]
    colors = ["#3aa655" if r["delta_vs_t00"] > (s if math.isfinite(s) else 0) else
              "#c0392b" if r["delta_vs_t00"] < -(s if math.isfinite(s) else 0) else "#8a8f98" for r in rows]
    ax.barh(ypos, [r["delta_vs_t00"] for r in rows], color=colors)
    if math.isfinite(s):
        ax.axvspan(-s, s, color="gray", alpha=0.18, label=f"±s = ±{s:.4f} (std của T00 qua {len(s2['choice']['t00_val_macro_f1'])} seed)")
        ax.legend(fontsize=8, loc="lower right")
    ax.axvline(0, color="black", lw=0.8)
    ax.set_yticks(ypos, [f"{r['exp_id']} [{r['axis']}] {r['diff']}"[:70] for r in rows], fontsize=8)
    ax.set(xlabel="Δ macro-F1 val so với T00 (seed 0)",
           title=f"Ablation công thức huấn luyện trên {s2['choice']['backbone']}\nxanh: Δ > s · đỏ: Δ < −s · xám: trong vùng nhiễu")
    ax.grid(alpha=0.3, axis="x")
    fig.tight_layout()
    path = lab.paths.figures / "training_ablation.png"
    fig.savefig(path, dpi=130)
    plt.close(fig)
    return path


# --------------------------------------------------------------------------- #
# report.md
# --------------------------------------------------------------------------- #
def md_table(df: pd.DataFrame, cols: list[str] | None = None, digits: int = 4) -> str:
    df = df[cols] if cols else df

    def cell(v):
        if v is None or (isinstance(v, float) and math.isnan(v)):
            return "–"
        if isinstance(v, (float, np.floating)):
            return f"{v:.{digits}f}"
        return str(v).replace("|", "/").replace("\n", " ")
    lines = ["| " + " | ".join(df.columns) + " |", "|" + "|".join("---" for _ in df.columns) + "|"]
    lines += ["| " + " | ".join(cell(v) for v in row) + " |" for row in df.itertuples(index=False)]
    return "\n".join(lines)


def write_report(lab, notebook_link: str = "(bổ sung link Google Colab trước khi nộp)") -> Path:
    T = tables(lab)
    d0, sn = lab.load("step0_data"), lab.load("step0_sanity")
    s1, s2, s3, s4 = (lab.load(k) for k in ("step1_backbones", "step2_training", "step3_inference", "step4_final"))
    errs, rb, ox = lab.load("step4_errors"), lab.load("bonus_robustness"), lab.load("bonus_onnx")
    sums = all_summaries(lab)
    seeds = list(lab.plan.seeds)
    L: list[str] = []
    w = L.append

    ef, et, e2, eu = (eval_summary(lab, t) for t in ("F01", "T00", "F02", "F01_uncal")) if s4 else (None,) * 4
    names = ef["classes"] if ef else CLASS_NAMES
    ia, isn = 0, 7
    noise = s2["choice"]["noise_std_t00"] if s2 else float("nan")
    bb = s2["choice"]["backbone"] if s2 else (s1["choice"]["chosen"] if s1 else "?")

    # ---------------------------------------------------------------- 1
    w("# Báo cáo Lab Day 2 — Backbone, công thức huấn luyện và suy luận trên DeepWeeds\n")
    w("**Sinh viên:** Nguyen Tuan Thanh · **MSSV:** 2A202602640 · Track 4, Ngày 2\n")
    w("> Mọi con số trong báo cáo lấy từ `results.xlsx` (sinh từ `logs/` và từ `eval.py` gốc chạy trên `predictions/`). "
      "Số ghi *trích dẫn* là của bài báo gốc hoặc bảng kết quả timm, không phải kết quả của bài này.\n")
    w("## 1. Tóm tắt\n")
    if s4:
        d = ef["macro_f1"]["mean"] - et["macro_f1"]["mean"]
        s_max = np.nanmax([ef["macro_f1"]["std"], et["macro_f1"]["std"]])
        ch = s4["inference"]
        w(f"- **Bài toán:** phân loại 9 lớp DeepWeeds (17.509 ảnh 256×256, `Negative` chiếm khoảng 52%), fold 0 chia sẵn, chỉ số chính macro-F1.")
        w(f"- **Đã làm:** {len(s1['rows'])} backbone cùng công thức nền (1 seed); {len([r for r in s2['rows'] if r['axis'] not in '-+'])} "
          f"thí nghiệm một-yếu-tố trên {len({r['axis'] for r in s2['rows'] if r['axis'] not in '-+'})} trục + "
          f"{len([r for r in s2['rows'] if r['axis'] == '+'])} kết hợp trên `{bb}`; {len([r for r in s3['rows'] if r.get('val_macro_f1') is not None]) - 1} "
          f"cấu hình suy luận ngoài mốc 1-view, có đo độ trễ p50/p95/p99; chung kết và mốc mỗi bên {len(seeds)} seed. Tổng {len(sums)} lần huấn luyện.")
        w(f"- **Cấu hình tốt nhất (F01):** `{bb}` + công thức `{s2['choice']['final_recipe_exp_id']}` "
          f"({next(r['diff'] for r in s2['rows'] if r['exp_id'] == s2['choice']['final_recipe_exp_id'])}) + suy luận {ch['offline_method']} "
          f"({ch['offline_desc']}) + temperature scaling.")
        w(f"- **Kết quả test ({len(seeds)} seed, test chạy một lần mỗi seed):** macro-F1 **{pm(ef['macro_f1']['mean'], ef['macro_f1']['std'])}**, "
          f"top-1 **{pm(ef['top1']['mean'], ef['top1']['std'])}**, ECE {pm(ef['ece']['mean'], ef['ece']['std'])}.")
        w(f"- **So với mốc T00 + I00:** macro-F1 test {pm(et['macro_f1']['mean'], et['macro_f1']['std'])}; Δ = {d:+.4f}, std lớn hơn của hai nhóm s = {s_max:.4f} → "
          + ("chênh lệch **vượt nhiễu**." if d > s_max else "chênh lệch **không vượt nhiễu**: không phân biệt được với mốc." if d > 0 else "chung kết **không tốt hơn** mốc trên test."))
        w(f"- **Hai lớp khó (recall test):** Chinee Apple {ef['recall']['mean'][ia] * 100:.1f}%, Snake Weed {ef['recall']['mean'][isn] * 100:.1f}% "
          f"(trích dẫn bài báo: {PAPER['chinee_apple']}% và {PAPER['snake_weed']}%).")
        w(f"- **Thời gian thực:** F02 (cùng mô hình, 1 view, FP32) có p95 batch-1 = {ch['i00_p95_ms']:.1f} ms trên {ch['latency_conditions']['gpu']} "
          f"({'đạt' if ch['i00_p95_ms'] <= 100 else 'không đạt'} ngân sách 100 ms), macro-F1 test {pm(e2['macro_f1']['mean'], e2['macro_f1']['std'])}.")
    else:
        w("_(Chưa có kết quả Bước 4.)_")
    w("")

    # ---------------------------------------------------------------- 2
    w("## 2. Dữ liệu và thiết lập\n")
    w(f"**Ngân sách:** cấu hình {lab.plan.epochs} epoch, batch {lab.plan.batch_size}; hạn nộp ngắn nên ưu tiên đủ thiết kế thí nghiệm. "
      + ("Số epoch ít hơn mức 10–15 đề xuất; kết quả chưa đại diện cho khả năng hội tụ tối đa. " if lab.plan.epochs < 10
         else "Số epoch nằm trong mức 10–15 đề xuất. " if lab.plan.epochs <= 15 else "Số epoch cao hơn mức 10–15 đề xuất. ")
      +
      "Các lần có cấu hình và seed giống hệt được tái sử dụng, ghi `reused_from` trong log, không coi là seed độc lập.\n")
    if d0:
        n = d0["n"]
        w(f"**Dataset và cách chia.** DeepWeeds (Olsen et al., 2019), fold 0 của tác giả, tải nguyên bản, không sửa/lọc/chia lại "
          f"(S1). Kết quả kiểm tra của `dataset.check_split` (`logs/step0_data.json`):\n")
        w(f"- Số ảnh: train **{n['train']}** / val **{n['val']}** / test **{n['test']}** "
          f"({d0['frac']['train'] * 100:.2f}% / {d0['frac']['val'] * 100:.2f}% / {d0['frac']['test'] * 100:.2f}%), tổng {d0['total']}.")
        w(f"- Giao theo tên file: train∩val = {d0['overlap']['train∩val']}, train∩test = {d0['overlap']['train∩test']}, "
          f"val∩test = {d0['overlap']['val∩test']}; hợp ba tập = **{d0['union']}** ảnh; file thiếu trong thư mục ảnh: {d0['missing_files']}.")
        w(f"- Định dạng ảnh: {d0['image_formats']} → mọi ảnh là RGB 256×256, nên val/test chỉ cần `CenterCrop(224)` và chuẩn hoá ImageNet.")
        w("")
        pcs = pd.DataFrame({"lớp": CLASS_NAMES, "train": d0["per_class"]["train"], "val": d0["per_class"]["val"],
                            "test": d0["per_class"]["test"], "tổng (đếm thật)": d0["class_total"],
                            "Table 1 bài báo (trích dẫn)": PAPER_TABLE1})
        w(md_table(pcs))
        w("")
        mis = d0.get("label_mismatch_vs_labels_csv", [])
        diff = [f"{CLASS_NAMES[i]} {a} (bài báo {b})" for i, (a, b) in enumerate(zip(d0["class_total"], PAPER_TABLE1)) if a != b]
        if d0["matches_paper_table1"]:
            w("Số đếm theo lớp **khớp** Table 1 của bài báo.")
        else:
            w("Số đếm theo lớp **lệch** Table 1 của bài báo ở: " + "; ".join(diff) + ". "
              + ("Nguyên nhân tìm được: " + "; ".join(
                  f"ảnh `{x['Filename']}` trong `{x['split']}_subset0.csv` mang nhãn {x['label_in_fold_csv']} ({CLASS_NAMES[x['label_in_fold_csv']]}) "
                  f"còn `labels.csv` ghi {x['label_in_labels_csv']} ({CLASS_NAMES[x['label_in_labels_csv']]})" for x in mis)
                 + ". Theo quy tắc S1, file CSV của fold được giữ nguyên, không sửa nhãn này (`eval.py` cũng đối chiếu `y_true` với file fold)."
                 if mis else "Không tìm thấy ảnh nào có nhãn trong file fold khác `labels.csv`."))
        p_neg = d0["class_total"][8] / d0["total"]
        w(f"Lớp nhiều nhất/lớp ít nhất = **{d0['imbalance_ratio']:.2f}** lần; `Negatives` chiếm {p_neg * 100:.1f}%, nên một mô hình luôn đoán "
          f"`Negatives` đã có top-1 ≈ {p_neg * 100:.0f}% nhưng macro-F1 chỉ ≈ {2 * p_neg / (1 + p_neg) / 9:.2f}. Vì vậy mọi lựa chọn dựa trên macro-F1 val.\n")
        w("![Phân bố lớp](figures/eda_class_distribution.png)\n")
        w("![Ảnh mẫu](figures/eda_samples.png)\n")
    w("**Quy tắc val/test.** Train chỉ để cập nhật trọng số; val để chọn backbone, công thức, phương pháp suy luận, checkpoint và nhiệt độ T; "
      "test chỉ mở ở Bước 4, đúng một lần cho mỗi seed (code ghi dấu `TEST_DONE.json` và từ chối chạy lại). Không gộp val vào train.\n")
    w("**Chỉ số.** Theo README mục 2.2, tính bằng `eval.compute_metrics`: macro-F1 (chính), top-1, balanced accuracy, P/R/F1 từng lớp, "
      "ECE 15 bin, mean ± std mẫu (ddof = 1).\n")
    w("**Công thức nền T00.** Trọng số ImageNet-1k (tag ghi trong sheet `Backbones`), head mới 9 lớp (khởi tạo N(0, 0,01²)), tinh chỉnh toàn bộ; "
      "train `RandomResizedCrop(224)` + lật ngang, val/test `CenterCrop(224)`; chuẩn hoá mean/std ImageNet; AdamW, LR backbone 1e-4 và head 1e-3, "
      "weight decay 0,05 (không áp dụng cho norm/bias/position embedding); warmup tuyến tính 1 epoch rồi cosine theo bước về 1% LR đỉnh; "
      f"cross-entropy; batch {lab.plan.batch_size}; {lab.plan.epochs} epoch; AMP; chọn checkpoint theo macro-F1 val (hòa lấy epoch sớm hơn).\n")
    if sums:
        v = sums[0]["versions"]
        w(f"**Phần cứng và thư viện.** {sums[0]['device']}; Python {v['python']}, torch {v['torch']}, torchvision {v['torchvision']}, timm {v['timm']}, "
          f"numpy {v['numpy']}, pandas {v['pandas']}. Seed quét sàng: {seeds[0]}; seed chung kết và mốc: {seeds}. "
          "`cudnn.benchmark=True` nên hai lần chạy cùng seed có thể lệch ở mức nhiễu số học; thứ tự batch, khởi tạo head và augmentation lặp lại được.\n")
    if sn:
        il = sn["initial_loss"]
        w("**Kiểm tra pipeline trước khi chạy thật** (`logs/step0_sanity.json`):\n")
        w(f"- Seed: hai DataLoader cùng seed cho cùng thứ tự file ({sn['seed_same_files']}) và cùng tensor sau augmentation ({sn['seed_same_tensors']}).")
        w(f"- Loss ban đầu (kỳ vọng −ln(1/9) = {LN9:.3f}): " + "; ".join(f"`{k.split('.')[0]}` {v:.3f}" for k, v in il.items()) + ".")
        o = sn["overfit"]
        w(f"- Overfit {o['n_images']} ảnh với `{o['backbone'].split('.')[0]}`: loss {o['first_loss']:.3f} → {o['final_loss']:.4f} sau {o['steps']} bước, "
          f"accuracy trên chính batch đó {o['final_acc_eval_mode'] * 100:.0f}%.")
        mc = sn["mode_check"]
        w(f"- `eval()`: đầu ra của một ảnh lệch {mc['eval_batch_dependence']:.1e} khi đổi các ảnh cùng batch; ở `train()` lệch {mc['train_batch_dependence']:.1e} "
          f"(BatchNorm dùng thống kê batch). Khi đóng băng backbone, số tầng BN còn ở train mode = {mc['frozen_bn_layers_in_train_mode']}.")
        w("- Ảnh sau augmentation và nhãn khớp nhau; CutMix/Mixup in kèm `lam` (hình dưới).\n")
        w("![Kiểm tra augmentation](figures/sanity_augmentation.png)\n")

    # ---------------------------------------------------------------- 3
    w("## 3. Kết quả so sánh backbone\n")
    if s1:
        w("Cùng công thức nền T00, cùng split, cùng seed 0; mỗi backbone 1 lần chạy (kết quả 1 seed).\n")
        w(md_table(T["Backbones"], ["exp_id", "backbone", "tag trọng số (timm)", "#tham số (M)", "GMAC", "macro-F1 val", "top-1 val",
                                    "thời gian train/epoch (s)", "độ trễ batch-1 p50 (ms)", "epoch tốt nhất"]))
        w("")
        w("![Backbone: F1 theo độ trễ và GMAC](figures/backbones_tradeoff.png)\n")
        ft = [r for r in s1["rows"] if r["init"] == "finetune"]
        ch = s1["choice"]
        best, worst = max(ft, key=lambda r: r["val_macro_f1"]), min(ft, key=lambda r: r["val_macro_f1"])
        chosen = next(r for r in ft if r["exp_id"] == ch["chosen_exp_id"])
        w(f"- **Tốt nhất về macro-F1 val:** {best['exp_id']} `{best['backbone'].split('.')[0]}` ({best['val_macro_f1']:.4f}); thấp nhất: "
          f"{worst['exp_id']} `{worst['backbone'].split('.')[0]}` ({worst['val_macro_f1']:.4f}); khoảng cách {best['val_macro_f1'] - worst['val_macro_f1']:.4f}."
          + (f" So với nhiễu seed đo ở Bước 2 (s = {noise:.4f}), " + ("khoảng cách giữa tốt nhất và thấp nhất vượt nhiễu, nhưng các backbone ở giữa "
             "cách nhau ít hơn s thì không xếp hạng được." if best['val_macro_f1'] - worst['val_macro_f1'] > noise else "cả khoảng này nằm trong nhiễu: không xếp hạng được các backbone.")
             if math.isfinite(noise) else ""))
        w(f"- **Chọn đi tiếp: {chosen['exp_id']} `{chosen['backbone']}`.** Quy tắc khai báo trước: {ch['rule']}. Các backbone gần tương đương với tốt nhất: "
          f"{', '.join(ch['near_best'])}; trong đó `{chosen['backbone'].split('.')[0]}` có macro-F1 val {chosen['val_macro_f1']:.4f} "
          f"(kém tốt nhất {best['val_macro_f1'] - chosen['val_macro_f1']:.4f}) và độ trễ p50 {chosen['latency_b1_p50_ms']:.1f} ms "
          f"(tốt nhất: {best['latency_b1_p50_ms']:.1f} ms).")
        notes = {r["exp_id"]: curve_notes(history(lab, r["exp_id"], r["seed"])) for r in s1["rows"]}
        fast = min(ft, key=lambda r: notes[r["exp_id"]]["epoch_reach_99pct"])
        w(f"- **Hội tụ:** `{fast['backbone'].split('.')[0]}` đạt 99% macro-F1 tốt nhất của chính nó sớm nhất (epoch {notes[fast['exp_id']]['epoch_reach_99pct']}). "
          + "Epoch đạt 99%: " + ", ".join(f"{r['exp_id']} = {notes[r['exp_id']]['epoch_reach_99pct']}" for r in ft) + ".")
        of = [r for r in s1["rows"] if notes[r["exp_id"]]["overfit_flag"]]
        w("- **Quá khớp:** " + ("; ".join(f"{r['exp_id']} có loss val thấp nhất ở epoch {notes[r['exp_id']]['val_loss_min_epoch']} "
          f"({notes[r['exp_id']]['val_loss_min']:.3f}) rồi tăng lên {notes[r['exp_id']]['val_loss_last']:.3f} ở epoch cuối trong khi loss train vẫn giảm" for r in of)
          + ". Macro-F1 val không giảm tương ứng, tức mô hình tự tin hơn trên các ảnh sai (ảnh hưởng hiệu chuẩn nhiều hơn độ chính xác)."
          if of else "không backbone nào có loss val tăng quá 10% so với mức thấp nhất trong khi loss train giảm (tiêu chí gắn cờ trong `make_results.curve_notes`)."))
        cited = [(r, IMAGENET_TOP1_CITED.get(r["weights_tag"])) for r in ft]
        if all(c is not None for _, c in cited) and len(cited) >= 3:
            a = pd.Series([r["val_macro_f1"] for r, _ in cited]).rank()
            b = pd.Series([c for _, c in cited]).rank()
            rho = float(np.corrcoef(a, b)[0, 1])
            w(f"- **So với thứ hạng ImageNet (trích dẫn bảng timm):** top-1 ImageNet của các tag là "
              + ", ".join(f"{r['exp_id']} {c}" for r, c in cited) + f". Tương quan hạng Spearman với macro-F1 val DeepWeeds: ρ = {rho:.2f} "
              f"({len(cited)} điểm, 1 seed nên chỉ mang tính mô tả).")
        g = np.array([r["gmacs"] for r in ft])
        lt = np.array([r["latency_b1_p50_ms"] for r in ft])
        tt = np.array([r["train_time_per_epoch_s"] for r in ft])
        w(f"- **GMAC có dự đoán được thời gian không?** Tương quan Pearson giữa GMAC và độ trễ batch-1: r = {np.corrcoef(g, lt)[0, 1]:.2f}; giữa GMAC và "
          f"thời gian train/epoch: r = {np.corrcoef(g, tt)[0, 1]:.2f}. "
          + "Ví dụ lệch: " + "; ".join(f"`{r['backbone'].split('.')[0]}` {r['gmacs']:.2f} GMAC → {r['latency_b1_p50_ms']:.1f} ms" for r in sorted(ft, key=lambda r: r['gmacs'])) + ".")
        x = next((r for r in s1["rows"] if r["init"] == "frozen"), None)
        if x:
            w(f"- **Bonus – linear probe DINOv2 đóng băng ({x['exp_id']}):** macro-F1 val {x['val_macro_f1']:.4f}, top-1 {x['val_top1']:.4f} với chỉ head được train, "
              f"so với CNN/transformer tinh chỉnh toàn bộ tốt nhất {best['val_macro_f1']:.4f} (chênh {x['val_macro_f1'] - best['val_macro_f1']:+.4f}). "
              "Lưu ý DINOv2 tiền huấn luyện trên LVD-142M, khác dữ liệu tiền huấn luyện của các backbone còn lại, nên đây không phải so sánh kiến trúc thuần.")
        w("")

    # ---------------------------------------------------------------- 4
    w("## 4. Kết quả công thức huấn luyện\n")
    if s2:
        c2 = s2["choice"]
        w(f"Backbone `{bb}`. Mỗi dòng T01… chỉ khác `T00` **đúng một yếu tố**, cùng seed 0, cùng số epoch. Nhiễu được đo bằng cách chạy `T00` với "
          f"{len(c2['t00_val_macro_f1'])} seed: macro-F1 val = {', '.join(f'{v:.4f}' for v in c2['t00_val_macro_f1'])} → mean {c2['t00_mean']:.4f}, "
          f"**s = {noise:.4f}** (std mẫu). Một yếu tố chỉ được gọi là có lợi/có hại khi |Δ| > s; ngược lại ghi \"không phân biệt được\". "
          "Vì các dòng ablation mới 1 seed, kể cả Δ > s cũng chỉ là bằng chứng sơ bộ.\n")
        w(md_table(T["Training"], ["exp_id", "trục thay đổi", "khác T00 ở điểm nào", "seed", "macro-F1 val", "top-1 val",
                                   "Δ macro-F1 so với T00 (seed 0)", "F1 val Chinee Apple", "F1 val Snake Weed", "ghi chú"]))
        w("")
        w("![Ablation](figures/training_ablation.png)\n")
        singles = [r for r in s2["rows"] if r["axis"] not in "-+"]
        thr = noise if math.isfinite(noise) else 0.0
        helps = [r for r in singles if r["delta_vs_t00"] > thr]
        hurts = [r for r in singles if r["delta_vs_t00"] < -thr]
        neutral = [r for r in singles if abs(r["delta_vs_t00"]) <= thr]
        fmt = lambda rs: ", ".join(f"{r['exp_id']} ({r['diff']}; Δ = {r['delta_vs_t00']:+.4f})" for r in rs) or "không có"  # noqa: E731
        w(f"- **Có lợi (Δ > s):** {fmt(helps)}.")
        w(f"- **Có hại (Δ < −s):** {fmt(hurts)}.")
        w(f"- **Không phân biệt được với nhiễu (|Δ| ≤ s):** {fmt(neutral)}.")
        for axis in sorted({r["axis"] for r in singles}):
            rs = [r for r in singles if r["axis"] == axis]
            b = max(rs, key=lambda r: r["val_macro_f1"])
            w(f"- **Trục {AXIS_NAMES[axis]}:** tốt nhất là {b['exp_id']} ({b['diff']}), Δ = {b['delta_vs_t00']:+.4f}"
              + (" (vượt nhiễu)." if b["delta_vs_t00"] > thr else " (không vượt nhiễu)."))
        base = s2["rows"][0]
        rare = max(singles, key=lambda r: r["f1_min_weed"])
        w(f"- **Lớp hiếm:** F1 val của lớp cỏ thấp nhất ở T00 là {base['f1_min_weed']:.4f}; cao nhất trong các ablation là {rare['exp_id']} ({rare['diff']}) "
          f"với {rare['f1_min_weed']:.4f} (Δ = {rare['f1_min_weed'] - base['f1_min_weed']:+.4f}); F1 `Negatives` tương ứng {base['f1_negatives']:.4f} → {rare['f1_negatives']:.4f}.")
        for r in [r for r in s2["rows"] if r["axis"] == "+"]:
            sd = r.get("sum_of_single_deltas", float("nan"))
            w(f"- **Kết hợp {r['exp_id']} ({r['diff']}):** Δ = {r['delta_vs_t00']:+.4f}, trong khi tổng Δ của từng yếu tố riêng lẻ là {sd:+.4f} → "
              + ("xấp xỉ cộng dồn" if abs(r["delta_vs_t00"] - sd) <= thr else "thấp hơn tổng (hiệu ứng chồng lấn/triệt tiêu một phần)" if r["delta_vs_t00"] < sd else "cao hơn tổng")
              + f" (so với s = {noise:.4f}).")
        fr = next(r for r in s2["rows"] if r["exp_id"] == c2["final_recipe_exp_id"])
        w(f"- **Công thức mang sang Bước 3–4:** {fr['exp_id']} ({fr['diff']}), macro-F1 val {fr['val_macro_f1']:.4f}, Δ = {fr['delta_vs_t00']:+.4f}. "
          f"Quy tắc: {c2['rule']}. Thiết kế là *tham lam theo trục nhưng chạy song song trên cùng nền T00* (không đổi nền giữa chừng), nên thứ tự trục "
          "không ảnh hưởng tới các Δ một-yếu-tố; việc chọn giá trị lớn nhất trong nhiều lần chạy 1 seed có thiên lệch lạc quan, và được kiểm lại bằng 3 seed ở Bước 4.")
        w("")

    # ---------------------------------------------------------------- 5
    w("## 5. Kết quả suy luận\n")
    if s3:
        c3 = s3["choice"]
        cond = c3["latency_conditions"]
        w(f"Mô hình: {c3['model']} (`{bb}`), không huấn luyện lại; mọi số đo trên **val**. Độ trễ đo bằng `benchmark.py`: {cond['warmup']} lần warmup bị bỏ, "
          f"`torch.cuda.synchronize()` trước và sau mỗi lần đo, {cond['iters']} lần đo, báo cáo p50/p95/p99; {cond['gpu']}, torch {cond['torch']}; "
          "đầu vào là tensor ngẫu nhiên đã nằm trên thiết bị (**không tính tiền xử lý**). TTA K view đo bằng K lượt forward liên tiếp cho một ảnh.\n")
        w(md_table(T["Inference"], ["exp_id", "phương pháp", "K (số view hoặc số mô hình)", "macro-F1 val", "top-1 val", "ECE val",
                                    "độ trễ p50 batch-1 (ms)", "độ trễ p95 batch-1 (ms)", "độ trễ p99 batch-1 (ms)",
                                    "chi phí tương đối so với I00", "Δ macro-F1 so với I00"]))
        w("")
        w("![Đánh đổi độ chính xác - độ trễ](figures/inference_tradeoff.png)\n")
        R = {r["exp_id"]: r for r in s3["rows"]}
        ok = lambda k: k in R and R[k].get("val_macro_f1") is not None  # noqa: E731
        thr = noise if math.isfinite(noise) else 0.0
        tag = lambda d: "vượt" if abs(d) > thr else "không vượt"  # noqa: E731
        if ok("I01"):
            w(f"- **TTA:** lật ngang (K=2) Δ = {R['I01']['delta_f1_vs_i00']:+.4f} với chi phí ×{R['I01']['rel_cost']:.2f}"
              + (f"; 10 crop Δ = {R['I02b']['delta_f1_vs_i00']:+.4f} với chi phí ×{R['I02b']['rel_cost']:.2f}" if ok("I02b") else "")
              + f". So với nhiễu seed s = {noise:.4f}, mức tăng của lật ngang {tag(R['I01']['delta_f1_vs_i00'])} nhiễu. "
              f"Phương pháp được chọn ({c3['offline_method']}) đổi {c3['offline_flips_vs_i00']['wrong_to_right']} ảnh val từ sai thành đúng và "
              f"{c3['offline_flips_vs_i00']['right_to_wrong']} ảnh từ đúng thành sai so với I00: TTA không phải độ chính xác miễn phí.")
        if ok("I03a") and ok("I01"):
            w(f"- **Gộp xác suất hay logit:** K=2: {R['I01']['val_macro_f1']:.4f} (xác suất) so với {R['I03a']['val_macro_f1']:.4f} (logit)"
              + (f"; K=10: {R['I02b']['val_macro_f1']:.4f} so với {R['I03b']['val_macro_f1']:.4f}" if ok("I03b") and ok("I02b") else "")
              + f". Chênh lệch tối đa {max(abs(R['I01']['val_macro_f1'] - R['I03a']['val_macro_f1']), abs(R['I02b']['val_macro_f1'] - R['I03b']['val_macro_f1']) if ok('I03b') and ok('I02b') else 0):.4f}, "
              "nhỏ so với nhiễu seed: không phân biệt được hai cách gộp.")
        res = [r for k, r in R.items() if k.startswith("I04")]
        if res:
            okres = [r for r in res if r.get("val_macro_f1") is not None]
            if okres:
                w("- **Độ phân giải kiểm tra (resize toàn ảnh):** " + "; ".join(f"{r['method'].split(' (')[0].replace('Độ phân giải kiểm tra ', '')}px: "
                  f"{r['val_macro_f1']:.4f} (×{r['rel_cost']:.2f})" for r in okres)
                  + f". Mốc I00 (cắt giữa 224 từ ảnh 256, tức vật thể to hơn 256/224 lần so với resize toàn ảnh về 224): {R['I00']['val_macro_f1']:.4f}. "
                  "Train bằng `RandomResizedCrop` làm vật thể lúc train trông to hơn lúc test (hiệu ứng FixRes), nên test ở độ phân giải cao hơn có thể có lợi; "
                  "số liệu trên cho biết điều đó có xảy ra ở đây hay không.")
            if len(okres) < len(res):
                w("- Một số độ phân giải không chạy được vì kiến trúc cố định đầu vào 224 (position embedding/cửa sổ attention).")
        ens = [R[k] for k in ("I05a", "I05b") if ok(k)]
        if ens:
            w("- **Ensemble:** " + "; ".join(f"{r['exp_id']} ({r['method']}): {r['val_macro_f1']:.4f}, Δ = {r['delta_f1_vs_i00']:+.4f}, chi phí ×{r['rel_cost']:.2f}" for r in ens)
              + ". I05b ghép các backbone huấn luyện bằng công thức nền nên so với I00 (mô hình chung kết) không phải so sánh một-yếu-tố.")
        if ok("I06a") and ok("I06b"):
            d = R["I06b"]["val_macro_f1"] - R["I06a"]["val_macro_f1"]
            w(f"- **EMA:** cùng lần chạy {R['I06b']['model']}, cùng epoch: trọng số thường {R['I06a']['val_macro_f1']:.4f}, trọng số EMA {R['I06b']['val_macro_f1']:.4f} "
              f"(Δ = {d:+.4f}, {tag(d)} nhiễu seed); không tốn thêm chi phí suy luận.")
        cal = c3["calibration_i00"]
        w(f"- **Hiệu chuẩn (I07):** T = {cal['T']:.3f} khớp trên val. ECE val {cal['ece_before']:.4f} → {cal['ece_after_insample']:.4f} (đo trên chính val nên lạc quan); "
          f"ước lượng cross-fit (khớp T trên nửa val này, đo trên nửa kia): {cal['ece_crossfit_before']:.4f} → {cal['ece_crossfit_after']:.4f}. "
          "Top-1 và macro-F1 không đổi vì chia logit cho T > 0 không đổi thứ tự lớp. Kiểm chứng độc lập trên test ở mục 6.")
        fu = c3["fuse"]
        if fu.get("has_batchnorm"):
            w(f"- **Gộp BatchNorm (I08a):** gộp {fu['pairs']} cặp conv–BN, logit lệch tối đa {fu['max_abs_diff_logits']:.1e}; "
              f"p50 batch-1 {R['I00']['lat_p50_ms']:.2f} → {R['I08a']['lat_p50_ms']:.2f} ms, macro-F1 val {R['I08a']['val_macro_f1']:.4f} (I00: {R['I00']['val_macro_f1']:.4f}).")
        else:
            w(f"- **Gộp BatchNorm (I08a):** không áp dụng được vì `{bb.split('.')[0]}` không có BatchNorm2d (dùng LayerNorm).")
        if ok("I08b") and ok("I08c"):
            w(f"- **AMP/FP16 ở batch 1:** FP32 p50 {R['I00']['lat_p50_ms']:.2f} ms, AMP {R['I08b']['lat_p50_ms']:.2f} ms, FP16 {R['I08c']['lat_p50_ms']:.2f} ms; "
              f"macro-F1 val lần lượt {R['I00']['val_macro_f1']:.4f} / {R['I08b']['val_macro_f1']:.4f} / {R['I08c']['val_macro_f1']:.4f}. "
              + ("AMP **chậm hơn** FP32 ở batch 1 trên máy này, đúng cảnh báo của slide trang 73." if R["I08b"]["lat_p50_ms"] > R["I00"]["lat_p50_ms"]
                 else "AMP nhanh hơn FP32 ở batch 1 trên máy này.") + " Số liệu batch 32 nằm ở sheet `Latency`.")
        w(f"- **Chọn cho chung kết (ngoại tuyến):** {c3['offline_method']} – {c3['offline_desc']} (macro-F1 val {c3['offline_val_macro_f1']:.4f}, "
          f"Δ = {c3['offline_delta_vs_i00']:+.4f} so với I00, p95 {c3['offline_p95_ms']:.1f} ms), rồi temperature scaling. Quy tắc: {c3['rule']}.")
        w(f"- **Cấu hình thời gian thực (F02):** 1 view FP32, p95 batch-1 = {c3['i00_p95_ms']:.2f} ms → {'đạt' if c3['i00_p95_ms'] <= 100 else 'KHÔNG đạt'} p95 ≤ 100 ms. "
          f"Biến thể nhanh nhất đo được là \"{c3['realtime_config']}\": p50 {c3['realtime_p50_ms']:.2f} / p95 {c3['realtime_p95_ms']:.2f} / p99 {c3['realtime_p99_ms']:.2f} ms "
          "(dự đoán F02 nộp trong `predictions/` được tính bằng FP32 không gộp BN; mức lệch xác suất của từng biến thể ghi ở cột ghi chú của sheet `Inference`). "
          "Nhận định của slide (TTA/ensemble hợp ngoại tuyến; trên robot dùng thứ không tốn thêm) được đối chiếu bằng cột chi phí tương đối ở bảng trên: "
          "TTA và ensemble nhân chi phí gần đúng theo K, còn EMA, temperature scaling, gộp BN không thêm lượt forward nào.")
        w("")

    # ---------------------------------------------------------------- 6
    w("## 6. Cấu hình tốt nhất\n")
    if s4:
        ch = s4["inference"]
        fcfg = lab.final_cfg(seeds[0])
        w("**Mô tả để tái lập.** "
          f"`{fcfg.backbone}` (timm, tiền huấn luyện ImageNet-1k), head 9 lớp; công thức = T00 + {{{', '.join(f'{k}={v}' for k, v in s2['choice']['final_overrides'].items()) or 'không đổi gì'}}}; "
          f"{fcfg.epochs} epoch, batch {fcfg.batch_size}, AdamW LR {fcfg.lr_backbone:g}/{fcfg.lr_head:g}, weight decay {fcfg.weight_decay}, warmup {fcfg.warmup_epochs:g} epoch + cosine, AMP; "
          f"checkpoint = epoch có macro-F1 val cao nhất; suy luận {ch['offline_method']} ({ch['offline_desc']}), view = {ch['offline_views']}, gộp theo {ch['offline_space']}; "
          "temperature scaling với T khớp trên val của từng seed (T = " + ", ".join(f"{x['T_offline']:.3f}" for x in s4["per_seed"]) + "). "
          f"Lệnh: `python train.py --set exp_id=F01 backbone={fcfg.backbone} seed=<k> " + " ".join(f"{k}={v}" for k, v in s2["choice"]["final_overrides"].items()) + "`.\n")
        w("**Bảng chung kết (test, tính bằng `eval.py score` từ `predictions/`).**\n")
        fin = T["Final"]
        w(md_table(fin[fin["seed"].astype(str).str.isdigit() | (fin["seed"] == "mean") | (fin["seed"] == "std")],
                   ["exp_id", "seed", "macro-F1 val", "macro-F1 test", "top-1 test", "balanced acc test", "ECE test"]))
        w("")
        d = ef["macro_f1"]["mean"] - et["macro_f1"]["mean"]
        s_max = float(np.nanmax([ef["macro_f1"]["std"], et["macro_f1"]["std"]]))
        w(f"- **F01 so với mốc T00:** macro-F1 test {pm(ef['macro_f1']['mean'], ef['macro_f1']['std'])} so với {pm(et['macro_f1']['mean'], et['macro_f1']['std'])}; "
          f"Δ = **{d:+.4f}**, s = {s_max:.4f} → " + ("Δ > s: cải thiện vượt nhiễu seed." if d > s_max else "0 < Δ ≤ s: không phân biệt được với mốc." if d > 0 else "Δ ≤ 0: không cải thiện.")
          + f" Top-1 test {ef['top1']['mean'] * 100:.2f}% so với {et['top1']['mean'] * 100:.2f}% (trích dẫn bài báo: ResNet-50 {PAPER['resnet50_acc']}%, Inception-v3 {PAPER['inception_acc']}%, "
          "huấn luyện khoảng 100 epoch với augmentation mạnh và trung bình 5 fold, nên chỉ để tham chiếu).")
        gap = abs(float(np.mean([x["F01"]["val_macro_f1"] for x in s4["per_seed"]])) - ef["macro_f1"]["mean"])
        w(f"- **Val so với test:** macro-F1 val trung bình {np.mean([x['F01']['val_macro_f1'] for x in s4['per_seed']]):.4f}, test {ef['macro_f1']['mean']:.4f}, chênh {gap:.4f} "
          + ("(≤ 0,02)." if gap <= 0.02 else "(> 0,02: val được dùng để chọn checkpoint, công thức và suy luận nên lạc quan hơn test)."))
        w(f"- **Hiệu chuẩn trên test:** ECE trước temperature scaling {pm(eu['ece']['mean'], eu['ece']['std'])}, sau {pm(ef['ece']['mean'], ef['ece']['std'])} "
          + ("→ T khớp trên val **giảm** ECE trên test." if ef["ece"]["mean"] < eu["ece"]["mean"] else "→ T khớp trên val **không giảm** ECE trên test."))
        w(f"- **Cấu hình thời gian thực F02** (cùng mô hình, 1 view + temperature scaling): macro-F1 test {pm(e2['macro_f1']['mean'], e2['macro_f1']['std'])}, "
          f"top-1 {pm(e2['top1']['mean'], e2['top1']['std'])}, p95 batch-1 {ch['i00_p95_ms']:.2f} ms (FP32).")
        w("")
        w("**Chỉ số theo lớp trên test (mean ± std qua seed).**\n")
        pc = T["PerClass"]
        both = pc[pc["exp_id"] == "F01"].merge(pc[pc["exp_id"] == "T00"], on="lớp", suffixes=(" F01", " T00"))
        tbl = pd.DataFrame({"lớp": both["lớp"], "số ảnh test": both["số ảnh test F01"],
                            "precision F01": [pm(a, b, 3) for a, b in zip(both["precision (mean) F01"], both["precision (std) F01"])],
                            "recall F01": [pm(a, b, 3) for a, b in zip(both["recall (mean) F01"], both["recall (std) F01"])],
                            "F1 F01": [pm(a, b, 3) for a, b in zip(both["F1 (mean) F01"], both["F1 (std) F01"])],
                            "F1 T00 (mốc)": [pm(a, b, 3) for a, b in zip(both["F1 (mean) T00"], both["F1 (std) T00"])]})
        w(md_table(tbl))
        w("")
        w(f"Hai lớp khó nhất theo bài báo: **Chinee Apple** precision {ef['precision']['mean'][ia]:.3f}, recall {ef['recall']['mean'][ia]:.3f}, F1 {ef['f1']['mean'][ia]:.3f}; "
          f"**Snake Weed** precision {ef['precision']['mean'][isn]:.3f}, recall {ef['recall']['mean'][isn]:.3f}, F1 {ef['f1']['mean'][isn]:.3f} "
          f"(trích dẫn bài báo, coi như recall: {PAPER['chinee_apple']}% và {PAPER['snake_weed']}%). "
          f"Lớp có F1 thấp nhất của F01: {names[int(np.argmin(ef['f1']['mean']))]} ({min(ef['f1']['mean']):.3f}).\n")
        w("**Ma trận nhầm lẫn và phân tích lỗi.**\n")
        w("![Ma trận nhầm lẫn test](figures/confusion_test.png)\n")
        if errs:
            top = errs["F01_top_confusions"]
            w(f"Các cặp bị nhầm nhiều nhất của F01 (cộng {len(seeds)} seed): "
              + "; ".join(f"{t['true']} → {t['pred']}: {t['count_sum_seeds']} lượt ({t['pct_of_true_class']:.1f}% số ảnh lớp thật)" for t in top[:6]) + ".")
            pair = [t for t in top if {t["true"], t["pred"]} == {"Chinee Apple", "Snake Weed"}]
            w(f"Cặp Chinee Apple ↔ Snake Weed: " + ("; ".join(f"{t['true']} → {t['pred']} {t['pct_of_true_class']:.1f}%" for t in pair) if pair else "không nằm trong 8 cặp nhầm nhiều nhất")
              + " (trích dẫn bài báo: 3,4% Chinee apple → Snake weed và 4,1% chiều ngược lại). "
              f"Với seed {seeds[0]}: {errs['n_wrong_seed']} ảnh test bị đoán sai, trong đó {errs['n_chinee_snake_confusions_seed']} ảnh thuộc cặp này.\n")
            w("![Ảnh test bị đoán sai](figures/errors_test.png)\n")
            if lab.load("bonus_gradcam"):
                w("![Grad-CAM trên ảnh bị đoán sai](figures/bonus_gradcam_errors.png)\n")
            w("**Giả thuyết nguyên nhân** *(cần rà lại bằng mắt trên hai hình trên trước khi nộp)*: (1) Chinee apple và Snake weed đều là tán lá nhỏ, xanh, "
              "chụp từ xa trên nền cỏ nên khác nhau chủ yếu ở kết cấu mịn, dễ mất khi cắt 224; (2) lỗi giữa một loài cỏ và `Negatives` hai chiều là loại lỗi phổ biến "
              "vì `Negatives` gồm mọi thực vật không phải mục tiêu, rất đa dạng, và nhiều ảnh cỏ dại chỉ chứa vật thể nhỏ hoặc bị che; "
              "(3) một số ảnh có thể gán nhãn theo vị trí chụp chứ loài mục tiêu chiếm rất ít điểm ảnh. Grad-CAM cho biết mô hình đang nhìn vào tán lá hay vào nền.\n")

    # ---------------------------------------------------------------- 7
    w("## 7. Kết luận và khuyến nghị\n")
    if s1 and s2 and s3 and s4:
        ft = [r for r in s1["rows"] if r["init"] == "finetune"]
        spread_bb = max(r["val_macro_f1"] for r in ft) - min(r["val_macro_f1"] for r in ft)
        gain_recipe = s2["choice"]["final_delta_vs_t00"]
        gain_inf = s3["choice"]["offline_delta_vs_i00"]
        d = ef["macro_f1"]["mean"] - et["macro_f1"]["mean"]
        s_max = float(np.nanmax([ef["macro_f1"]["std"], et["macro_f1"]["std"]]))
        w(f"1. **Cấu hình tốt nhất:** F01 = `{bb}` + công thức {s2['choice']['final_recipe_exp_id']} + {s3['choice']['offline_method']} + temperature scaling; macro-F1 test "
          f"{pm(ef['macro_f1']['mean'], ef['macro_f1']['std'])}, hơn mốc {d:+.4f} (s = {s_max:.4f}: "
          + ("vượt nhiễu" if d > s_max else "không vượt nhiễu") + ").")
        parts = sorted([("chọn backbone (khoảng cách tốt nhất − thấp nhất, val, 1 seed)", spread_bb),
                        ("công thức huấn luyện (công thức chung kết − T00, val, seed 0)", gain_recipe),
                        ("suy luận (phương pháp được chọn − I00, val)", gain_inf)], key=lambda t: -t[1])
        w("2. **Yếu tố đóng góp nhiều nhất** (theo macro-F1 val): " + "; ".join(f"{n}: {v:+.4f}" for n, v in parts)
          + f". Nhiễu seed s = {noise:.4f}: các mức nhỏ hơn s không nên coi là đóng góp thật. Lưu ý khoảng cách backbone là giữa mạng tốt nhất và mạng thấp nhất, "
          "không phải mức tăng so với mốc ResNet.")
        c3 = s3["choice"]
        w(f"3. **Triển khai trên robot (30–100 ms/khung):** dùng F02 – `{bb}` 1 view FP32, p95 = {c3['i00_p95_ms']:.1f} ms (biến thể nhanh nhất \"{c3['realtime_config']}\": {c3['realtime_p95_ms']:.1f} ms) trên {c3['latency_conditions']['gpu']}, "
          f"macro-F1 test {pm(e2['macro_f1']['mean'], e2['macro_f1']['std'])}; giữ temperature scaling (không tốn chi phí). "
          f"{c3['offline_method']} có p95 = {c3['offline_p95_ms']:.1f} ms" + (" – vẫn trong ngân sách trên GPU này nhưng" if c3["offline_p95_ms"] <= 100 else " – vượt ngân sách, và")
          + f" chỉ thêm {c3['offline_delta_vs_i00']:+.4f} macro-F1 val, nên dành cho xử lý ngoại tuyến. Độ trễ đo trên GPU máy chủ và không gồm tiền xử lý; trên phần cứng nhúng "
          "(bài báo: Jetson TX2, ResNet-50 53–180 ms, trích dẫn) cần đo lại, và khi đó các backbone nhẹ ở sheet `Backbones` là phương án dự phòng.")
        w("")

    # ---------------------------------------------------------------- 8
    w("## 8. Hạn chế và việc tiếp theo\n")
    w(f"- **Số seed:** quét backbone và ablation chỉ 1 seed (seed {seeds[0]}); chỉ T00 và F01 có {len(seeds)} seed. Nhiễu s ước lượng từ {len(seeds)} seed của T00 nên bản thân nó cũng không chắc.")
    w("- **Một fold:** chỉ fold 0. Dữ liệu chia ngẫu nhiên, không theo địa điểm chụp, nên ảnh cùng địa điểm/cùng đợt chụp nằm ở cả train và test; điểm test có thể **lạc quan** so với địa điểm mới.")
    w("- **Thiên lệch chọn lựa:** công thức và phương pháp suy luận được chọn là giá trị lớn nhất trong nhiều lần đo trên cùng tập val; mức tăng trên val vì thế lạc quan (so sánh val–test ở mục 6).")
    w(f"- **Ngân sách GPU:** {lab.plan.epochs} epoch (bài báo khoảng 100), không dò LR riêng cho transformer, ablation chỉ trên một backbone, mỗi trục vài giá trị với siêu tham số mặc định "
      "(ví dụ CutMix α = 1 với xác suất 0,5; focal γ = 2; EMA 0,998) nên \"không giúp\" ở đây không có nghĩa là kỹ thuật đó vô ích.")
    w("- **Độ trễ:** đo forward trên GPU máy chủ, đầu vào ngẫu nhiên, không tính giải mã ảnh/tiền xử lý/chép dữ liệu; không đại diện cho phần cứng nhúng.")
    w("- **Tái lập:** `cudnn.benchmark=True`, không ép thuật toán tất định; chạy lại cùng seed cho kết quả cùng mức, không trùng từng chữ số.")
    if rb:
        r_f = [r for r in rb["rows"] if r["model"].startswith("F01")]
        clean = next(r for r in r_f if r["corruption"].startswith("sạch"))
        worst = min(r_f, key=lambda r: r["val_macro_f1"])
        bad_ts = [r for r in r_f if r["ece_after_ts"] > r["ece_before_ts"]]
        w(f"- **Lệch phân phối (bonus, trên val bị làm hỏng):** macro-F1 của F01 từ {clean['val_macro_f1']:.4f} (sạch) xuống thấp nhất {worst['val_macro_f1']:.4f} với \"{worst['corruption']}\"; "
          f"ECE sau temperature scaling (T từ val sạch) ở điều kiện đó là {worst['ece_after_ts']:.4f} so với {clean['ece_after_ts']:.4f} trên val sạch. "
          + (f"T khớp trên val sạch làm ECE **xấu đi** ở: {', '.join(r['corruption'] for r in bad_ts)}. " if bad_ts else "T khớp trên val sạch vẫn giảm ECE ở mọi điều kiện đã thử. ")
          + "Kết luận: T (và cả độ chính xác) khớp trên val cùng phân phối không đảm bảo khi ánh sáng/mùa/camera thay đổi.")
        w("")
        w(md_table(pd.DataFrame(rb["rows"])[["model", "corruption", "val_macro_f1", "val_top1", "ece_before_ts", "ece_after_ts", "min_weed_recall"]]))
        w("")
        w("![Các kiểu lệch phân phối tự tạo](figures/bonus_corruptions.png)\n")
    if ox and ox.get("rows"):
        w("- **ONNX (bonus):** " + "; ".join(f"{r['runtime']}: p50 {r['p50']:.2f} ms, p95 {r['p95']:.2f} ms" + (f", lệch logit tối đa {r['max_abs_diff_vs_pytorch']:.1e}" if r.get("max_abs_diff_vs_pytorch") else "")
                                            if "p50" in r else f"{r['runtime']}: không chạy được ({r.get('error', '')[:80]})" for r in ox["rows"])
          + ". ONNX Runtime đo gồm cả chép đầu vào numpy vào session.")
    w("- **Việc tiếp theo:** chạy đủ 5 fold (hoặc chia theo địa điểm) để có ước lượng không lạc quan; thêm seed cho các ablation sát ngưỡng nhiễu; "
      "huấn luyện dài hơn với augmentation hình học mạnh như bài báo (xoay 360°); chưng cất mô hình tốt nhất sang mạng nhẹ; đo độ trễ trên thiết bị nhúng thật.")
    w("")

    # ---------------------------------------------------------------- 9
    w("## 9. Phụ lục\n")
    w(f"- Notebook chạy lại được: {notebook_link}")
    w("- Code: `code/` (một hàm `train.run(Config)` cho mọi thí nghiệm; `selftest.py` kiểm tra focal γ=0, label smoothing ε=0, CutMix, nhóm tham số, EMA, lịch LR, gộp BN, temperature scaling, đo độ trễ). `eval.py` dùng nguyên bản.")
    w("- Truy vết: mỗi `exp_id` có `logs/<exp_id>/seed<k>/{config.json, history.csv, summary.json}` và ảnh `curves/<exp_id>_*.png`.\n")
    if sums:
        ap = pd.DataFrame([{"exp_id": s["exp_id"], "seed": s["seed"], "backbone (tag)": s["weights_tag"], "khác công thức nền": s["diff"],
                            "epoch tốt nhất": s["best_epoch"], "macro-F1 val": s["val_macro_f1"], "ảnh đường cong": f"curves/{s['curve']}"} for s in sums])
        w(md_table(ap))
        w("")
    path = lab.paths.sub / "report.md"
    path.write_text("\n".join(L), encoding="utf-8")
    return path


# --------------------------------------------------------------------------- #
# Kiểm tra nhất quán trước khi nộp
# --------------------------------------------------------------------------- #
def check_consistency(lab) -> dict:
    """exp_id nào cũng có ảnh trong curves/; số trong sheet Final khớp eval_out/ (kết quả của eval.py)."""
    out = {"missing_curves": [], "final_mismatch": []}
    for s in all_summaries(lab):
        if not (lab.paths.curves / s["curve"]).exists():
            out["missing_curves"].append(f"{s['exp_id']} seed {s['seed']}")
    xlsx = lab.paths.sub / "results.xlsx"
    if xlsx.exists() and lab.load("step4_final"):
        fin = pd.read_excel(xlsx, sheet_name="Final")
        for tag in ("T00", "F01", "F02"):
            es = eval_summary(lab, tag)
            row = fin[(fin["exp_id"] == tag) & (fin["seed"] == "mean")].iloc[0]
            for col, key in (("macro-F1 test", "macro_f1"), ("top-1 test", "top1"), ("ECE test", "ece")):
                if abs(float(row[col]) - es[key]["mean"]) > 1e-9:
                    out["final_mismatch"].append(f"{tag} {col}")
    out["n_runs"] = len(all_summaries(lab))
    out["n_curves"] = len(list(lab.paths.curves.glob("*.png")))
    out["ok"] = not out["missing_curves"] and not out["final_mismatch"]
    print("Kiểm tra nhất quán:", "ĐẠT" if out["ok"] else "CÓ LỖI", out)
    return out


def update_readme(lab, notebook_link: str | None = None) -> None:
    """Ghi phiên bản thư viện, phần cứng, seed và link notebook vào README.md của bài nộp (giữa các dấu RUN_INFO)."""
    import re

    path = lab.paths.sub / "README.md"
    sums = all_summaries(lab)
    if not path.exists() or not sums:
        return
    v = sums[0]["versions"]
    devices = sorted({s["device"] for s in sums})
    block = "\n".join([
        "| Mục | Giá trị |", "|---|---|",
        f"| GPU/thiết bị huấn luyện | {', '.join(devices)} |",
        *[f"| {k} | {val} |" for k, val in v.items()],
        f"| Seed | quét sàng: {lab.plan.seeds[0]}; mốc T00 và chung kết F01: {list(lab.plan.seeds)} |",
        f"| Epoch / batch | {lab.plan.epochs} / {lab.plan.batch_size} |",
        f"| Số lần huấn luyện | {len(sums)} (tổng {sum(s['total_time_s'] for s in sums) / 3600:.2f} giờ GPU) |",
        f"| Công cụ đếm GMAC | {sums[0]['gmac_tool']} |"])
    text = path.read_text(encoding="utf-8")
    text = re.sub(r"<!-- RUN_INFO -->.*?<!-- /RUN_INFO -->", lambda _: f"<!-- RUN_INFO -->\n{block}\n<!-- /RUN_INFO -->", text, flags=re.S)
    if notebook_link:
        text = re.sub(r"<!-- NOTEBOOK_LINK -->.*?<!-- /NOTEBOOK_LINK -->",
                      lambda _: f"<!-- NOTEBOOK_LINK -->{notebook_link}<!-- /NOTEBOOK_LINK -->", text, flags=re.S)
    path.write_text(text, encoding="utf-8")


def build_all(lab, notebook_link: str | None = None) -> dict:
    if lab.load("step4_final"):
        lab.refresh_scores()
    backbone_figure(lab)
    ablation_figure(lab)
    xlsx = build_xlsx(lab)
    rep = write_report(lab, notebook_link) if notebook_link else write_report(lab)
    update_readme(lab, notebook_link)
    return {"xlsx": str(xlsx), "report": str(rep), **check_consistency(lab)}

# Tham khảo implementation: picuisme/K4-Track4-Day2-Deeplearning-Advance @ 941d9fb.
# Bản cho Nguyen Tuan Thanh: sửa optimizer/resume/report/Colab; số liệu phải chạy riêng.
